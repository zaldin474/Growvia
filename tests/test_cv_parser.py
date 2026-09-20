from pathlib import Path

import pytest

from io import BytesIO
from docx import Document

from cv_analysis.step01_cv_parser import extract_cv_text


# -----------------------------------------
# TXT
# -----------------------------------------

def test_txt_cv():

    file_bytes = b"""
    John Doe

    Python Developer
    john@example.com
    """

    result = extract_cv_text(file_bytes,  "cv.txt")

    assert "John Doe" in result
    
    assert "Python Developer" in result


# -----------------------------------------
# Unsupported format
# -----------------------------------------

def test_unsupported_file():

    with pytest.raises(ValueError,  match="Unsupported CV format"):

        extract_cv_text( b"fake data", "cv.exe")


# -----------------------------------------
# Empty CV
# -----------------------------------------

def test_empty_cv():

    with pytest.raises(ValueError,  match="No readable text"):

        extract_cv_text( b"      ", "cv.txt")


# -----------------------------------------
# REAL TEST PDF
# -----------------------------------------

def test_real_pdf():

    growvia_dir = ( Path(__file__).resolve().parents[1] )

    cv_path = (
        growvia_dir
        / "test_cv.pdf"
    )

    file_bytes = cv_path.read_bytes()

    result = extract_cv_text(file_bytes,  cv_path.name)

    assert len(result) > 100

    assert ("ZEINEDDIN ZIDAN" in  result)

    assert ("TECHNICAL SKILLS" in  result)
    
    
def test_corrupted_pdf():

    with pytest.raises(
        ValueError,
        match="could not be parsed"
    ):

        extract_cv_text(
            b"this is not a pdf",
            "resume.pdf"
        )


def test_corrupted_docx():

    with pytest.raises(
        ValueError,
        match="could not be parsed"
    ):

        extract_cv_text(
            b"this is not a docx",
            "resume.docx"
        )


def test_docx_table_content():

    document = Document()

    document.add_paragraph(
        "John Doe"
    )

    table = document.add_table(
        rows=2,
        cols=1
    )

    table.cell(
        0,
        0
    ).text = "TECHNICAL SKILLS"

    table.cell(
        1,
        0
    ).text = "Python, Java, SQL"

    buffer = BytesIO()

    document.save(
        buffer
    )

    result = extract_cv_text(
        buffer.getvalue(),
        "resume.docx"
    )

    assert (
        "TECHNICAL SKILLS"
        in result
    )

    assert (
        "Python, Java, SQL"
        in result
    )